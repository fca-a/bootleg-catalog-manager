"""Layout-based XLSX exporter.

Consumes a :class:`ResolvedExportDocument` and produces an XLSX file.

Public API
----------
- :func:`export_layout_xlsx` — the single entry point.
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from bootleg_catalog.core.layout_service import (
    GridCell,
    ScreenshotLayout,
    compute_grid,
)
from bootleg_catalog.models.resolved_export import ResolvedExportDocument


class LayoutXLSXExportError(Exception):
    """Raised when the layout XLSX export fails."""


# ---------------------------------------------------------------------------
# Page setup (best effort)
# ---------------------------------------------------------------------------

# Excel paperSize codes (ECMA-376) for the supported page sizes.
_EXCEL_PAPER_SIZE: dict[str, int] = {"A4": 9, "Letter": 1, "A3": 8}
_MM_PER_INCH = 25.4


def _apply_page_setup(ws, page_config) -> None:
    """Seed the worksheet's print setup from *page_config*.

    Excel has no concept of pages until printing, so this only sets
    orientation, paper size, and margins — the content itself is NOT
    paginated here.
    """
    ws.page_setup.orientation = page_config.orientation  # portrait/landscape
    ws.page_setup.paperSize = _EXCEL_PAPER_SIZE.get(page_config.page_size, 9)

    from openpyxl.worksheet.page import PageMargins

    top, right, bottom, left = page_config.effective_margins_mm()
    ws.page_margins = PageMargins(
        left=left / _MM_PER_INCH,
        right=right / _MM_PER_INCH,
        top=top / _MM_PER_INCH,
        bottom=bottom / _MM_PER_INCH,
    )


# ---------------------------------------------------------------------------
# Validation helpers
# ---------------------------------------------------------------------------

def _validate_screenshots(
    screenshot_paths_by_recording: dict[str, list[str]],
) -> None:
    """Raise LayoutXLSXExportError if any screenshot file is missing."""
    missing: list[str] = []
    for paths in screenshot_paths_by_recording.values():
        for p in paths:
            if not Path(p).is_file():
                missing.append(p)
    if missing:
        names = "\n".join(f"  {p}" for p in missing)
        raise LayoutXLSXExportError(
            f"Screenshot file(s) not found ({len(missing)}):\n{names}"
        )


def _validate_destination(path: Path) -> None:
    """Raise LayoutXLSXExportError if destination already exists."""
    if path.exists():
        raise LayoutXLSXExportError(
            f"File already exists: {path}\n"
            "Delete it first or choose a different name."
        )


# ---------------------------------------------------------------------------
# Image helpers (duplicated from layout_pdf_exporter)
# These are identical to the PDF exporter's helpers.  We duplicate them
# to keep each exporter self-contained with no cross-module coupling.
# ---------------------------------------------------------------------------

def _image_dimensions(path: Path) -> tuple[int, int]:
    """Return (width, height) in pixels of an image file."""
    from PIL import Image

    with Image.open(path) as img:
        return img.size


def _fit_image(
    img_w: int, img_h: int, max_w: float, max_h: float
) -> tuple[float, float]:
    """Scale image to fit within max_w × max_h, preserving aspect ratio."""
    if img_w <= 0 or img_h <= 0:
        return max_w, max_h
    scale = min(max_w / img_w, max_h / img_h)
    if scale >= 1.0:
        return float(img_w), float(img_h)
    return img_w * scale, img_h * scale


def _compute_row_max_heights(
    cells: list[GridCell],
    cell_w: float,
    max_h: float,
) -> dict[int, float]:
    """Compute the maximum scaled height per row across all cells."""
    row_heights: dict[int, float] = {}
    for cell in cells:
        img_w, img_h = _image_dimensions(Path(cell.screenshot_path))
        _, scaled_h = _fit_image(img_w, img_h, cell_w, max_h)
        prev = row_heights.get(cell.row, 0.0)
        if scaled_h > prev:
            row_heights[cell.row] = scaled_h
    return row_heights


# ---------------------------------------------------------------------------
# Style helpers
# ---------------------------------------------------------------------------

def _make_font(style=None, *, bold: bool = False, italic: bool = False,
               size: int = 10, color: str | None = None):
    """Create an openpyxl Font from an ElementStyle or explicit kwargs."""
    from openpyxl.styles import Font

    if style is not None:
        bold = style.bold
        italic = style.italic
        size = style.font_size if style.font_size is not None else 10
        color = style.color

    font_kwargs: dict = dict(bold=bold, italic=italic, size=size)
    if color:
        # Expect hex like "#FF0000" — openpyxl wants "FF0000"
        hex_color = color.lstrip("#")
        if len(hex_color) == 6:
            font_kwargs["color"] = hex_color
    return Font(**font_kwargs)


_DEFAULT_FONT = _make_font(size=10)
_SHOW_HEADER_FONT = _make_font(size=12, bold=True)

# Light gray fill for show headers
_SHOW_HEADER_FILL = None  # Lazy-init to avoid top-level import


def _get_show_header_fill():
    global _SHOW_HEADER_FILL
    if _SHOW_HEADER_FILL is None:
        from openpyxl.styles import PatternFill
        _SHOW_HEADER_FILL = PatternFill(
            start_color="EEEEEE", end_color="EEEEEE", fill_type="solid",
        )
    return _SHOW_HEADER_FILL


# ---------------------------------------------------------------------------
# Main export function
# ---------------------------------------------------------------------------

# Column merge range for document header
_HEADER_MERGE_COLS = "A:F"


def export_layout_xlsx(
    document: ResolvedExportDocument,
    output_path: str | Path,
    *,
    screenshot_paths_by_recording: dict[str, list[str]] | None = None,
) -> Path:
    """Generate an XLSX from a fully-resolved export document.

    All shows are placed on a single worksheet, grouped visually with
    styled section headers.

    Parameters
    ----------
    document:
        The resolved document (already resolved by LayoutResolver).
    output_path:
        Destination for the XLSX file.
    screenshot_paths_by_recording:
        Optional mapping of ``recording.source_path`` → list of
        screenshot file paths.  When ``None``, no screenshots are
        rendered and no filesystem access occurs.

    Returns
    -------
    Path
        Absolute path of the created XLSX.

    Raises
    ------
    LayoutXLSXExportError
        If validation fails or generation errors.

    Validation order
    ----------------
    1. Destination path check.
    2. Screenshot file existence check (only if screenshots provided).
    """
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.utils import get_column_letter

    output = Path(output_path)

    # --- Validate ---
    _validate_destination(output)
    if screenshot_paths_by_recording:
        _validate_screenshots(screenshot_paths_by_recording)

    screenshot_config = document.layout.screenshot_config

    wb = Workbook()
    ws = wb.active
    ws.title = "Catalog"

    # Best-effort print setup from the layout's page config.
    _apply_page_setup(ws, document.layout.page_config)

    # Set column A width for readability
    ws.column_dimensions["A"].width = 40

    row = 1

    # --- Document header ---
    ws.merge_cells(f"A{row}:{_HEADER_MERGE_COLS.split(':')[1]}{row}")
    cell = ws.cell(row=row, column=1, value="Bootleg Catalog Export")
    cell.font = _make_font(size=14, bold=True)
    row += 2  # blank row after header

    # --- Show sections ---
    for section_idx, section in enumerate(document.show_sections):
        # Blank row between sections (not before the first)
        if section_idx > 0:
            row += 1

        # --- Show header lines ---
        for line in section.header_lines:
            text = "".join(elem.text for elem in line.elements)
            cell = ws.cell(row=row, column=1, value=text)
            cell.font = _SHOW_HEADER_FONT
            cell.fill = _get_show_header_fill()
            # Merge across a reasonable range so the header spans the sheet
            ws.merge_cells(f"A{row}:{_HEADER_MERGE_COLS.split(':')[1]}{row}")
            row += 1

        row += 1  # blank row after header

        # --- Recordings ---
        for rec_idx, resolved_rec in enumerate(section.recordings):
            if rec_idx > 0:
                row += 1  # blank row between recordings

            # Render body lines
            for line in resolved_rec.lines:
                text = "".join(elem.text for elem in line.elements)
                cell = ws.cell(row=row, column=1, value=text)

                # Apply dominant style (last element with a style)
                last_style = None
                for elem in line.elements:
                    if elem.style is not None:
                        last_style = elem.style
                cell.font = _make_font(last_style) if last_style else _DEFAULT_FONT
                row += 1

            # --- Screenshots ---
            if (
                screenshot_paths_by_recording is not None
                and screenshot_config.enabled
            ):
                src_path = resolved_rec.recording.source_path
                screenshots = screenshot_paths_by_recording.get(src_path, [])
                if screenshots:
                    row = _embed_screenshots(
                        ws, screenshots, row, screenshot_config,
                    )
                    row += 1  # blank row after screenshots

    # --- Save via temp file ---
    try:
        output.parent.mkdir(parents=True, exist_ok=True)
        tmp_fd, tmp_path = tempfile.mkstemp(
            suffix=".xlsx", dir=str(output.parent),
        )
        try:
            os.close(tmp_fd)
            wb.save(tmp_path)
            Path(tmp_path).replace(output)
        except Exception:
            try:
                Path(tmp_path).unlink(missing_ok=True)
            except OSError:
                pass
            raise
    except LayoutXLSXExportError:
        raise
    except Exception as exc:
        raise LayoutXLSXExportError(
            f"Failed to generate XLSX: {exc}"
        ) from exc

    return output.resolve()


# ---------------------------------------------------------------------------
# Screenshot embedding
# ---------------------------------------------------------------------------

def _embed_screenshots(
    ws,
    screenshot_paths: list[str],
    start_row: int,
    screenshot_config,
) -> int:
    """Embed screenshots into the worksheet and return the next free row."""
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.utils import get_column_letter

    photos_per_row = (
        3 if screenshot_config.layout == "grid" else len(screenshot_paths)
    )
    layout = ScreenshotLayout(photos_per_row=photos_per_row)
    cells = compute_grid(screenshot_paths, layout)
    if not cells:
        return start_row

    max_w = float(screenshot_config.max_width_px)
    max_h = float(screenshot_config.max_height_px)

    # Pre-compute scaled dimensions per row for row height
    row_heights = _compute_row_max_heights(cells, max_w, max_h)

    for cell in cells:
        img_path = Path(cell.screenshot_path)
        img_w, img_h = _image_dimensions(img_path)
        scaled_w, scaled_h = _fit_image(img_w, img_h, max_w, max_h)

        r = start_row + cell.row
        c = cell.col + 1  # 1-based

        img = XlImage(str(img_path))
        img.width = int(scaled_w)
        img.height = int(scaled_h)
        cell_ref = f"{get_column_letter(c)}{r}"
        ws.add_image(img, cell_ref)

    # Set row heights for image rows
    for row_idx, row_h_px in row_heights.items():
        r = start_row + row_idx
        # Convert px to points (1 px ≈ 0.75 pt at 96 dpi) + padding
        ws.row_dimensions[r].height = int(row_h_px * 0.75 + 10)

    # Return next free row
    max_row = max(cell.row for cell in cells)
    return start_row + max_row + 1
